import pandas as pd
import time
import os
import json
import datetime
import re
import threading
from flask import Flask, render_template, request, Response, send_file, jsonify

from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.chrome.options import Options
from selenium.common.exceptions import TimeoutException, WebDriverException, StaleElementReferenceException
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

app = Flask(__name__)

# Folders
UPLOAD_FOLDER = "uploads"
RESULT_FOLDER = "results"
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(RESULT_FOLDER, exist_ok=True)

# Real LASIK URL
LAPAKASIK_URL = "https://jmoantrian.bpjsketenagakerjaan.go.id/?source=e419a6aed6c50fefd9182774c25450b333de8d5e29169de6018bd1abb1c8f89b"

driver = None
wait = None
is_running = False
stop_flag = False

# Counter
sukses_count = 0
gagal_count = 0

# Data sementara
processed_rows = []
current_output_path = None
excel_path = None

# Deteksi kolom otomatis
detected_columns = {
    'nik': -1,
    'kpj': -1,
    'nama': -1,
    'kelurahan': -1
}
all_headers = []

# Auto mode - tidak perlu konfirmasi manual
AUTO_MODE = True

class ProcessedRow:
    def __init__(self, row_number, all_columns_data, kpj, nama, nik, status, saldo_jht, status_jmo, nama_pt):
        self.row_number = row_number
        self.all_columns_data = all_columns_data
        self.kpj = kpj
        self.nama = nama
        self.nik = nik
        self.status = status
        self.saldo_jht = saldo_jht
        self.status_jmo = status_jmo
        self.nama_pt = nama_pt

def detect_columns(df):
    """Deteksi kolom secara otomatis"""
    global detected_columns, all_headers
    
    detected_columns = {
        'nik': -1,
        'kpj': -1,
        'nama': -1,
        'kelurahan': -1
    }
    all_headers = list(df.columns)
    
    print("🔍 Mendeteksi kolom secara otomatis...")
    
    for idx, header in enumerate(all_headers):
        header_lower = str(header).lower()
        
        if detected_columns['nik'] == -1 and any(keyword in header_lower for keyword in ['nik', 'no ktp', 'noktp', 'e-ktp', 'ektp', 'ktp', 'identitas']):
            detected_columns['nik'] = idx
            print(f"✅ Kolom NIK terdeteksi: '{header}' (Kolom {idx + 1})")
        
        if detected_columns['kpj'] == -1 and any(keyword in header_lower for keyword in ['kpj', 'no kpj', 'nokpj', 'peserta', 'no peserta']):
            detected_columns['kpj'] = idx
            print(f"✅ Kolom KPJ terdeteksi: '{header}' (Kolom {idx + 1})")
        
        if detected_columns['nama'] == -1 and any(keyword in header_lower for keyword in ['nama', 'name', 'nm', 'peserta', 'karyawan']):
            detected_columns['nama'] = idx
            print(f"✅ Kolom NAMA terdeteksi: '{header}' (Kolom {idx + 1})")
        
        if detected_columns['kelurahan'] == -1 and any(keyword in header_lower for keyword in ['kelurahan', 'desa', 'kecamatan', 'alamat']):
            detected_columns['kelurahan'] = idx
            print(f"✅ Kolom KELURAHAN terdeteksi: '{header}' (Kolom {idx + 1})")
    
    if detected_columns['nik'] == -1 or detected_columns['kpj'] == -1:
        print("⚠️ Melakukan fallback detection berdasarkan data...")
        sample_row = df.iloc[0] if len(df) > 0 else None
        
        if sample_row is not None:
            for idx, value in enumerate(sample_row):
                value_str = str(value).strip() if pd.notna(value) else ""
                
                if detected_columns['nik'] == -1 and len(value_str) >= 16 and value_str.isdigit():
                    detected_columns['nik'] = idx
                    print(f"✅ Kolom NIK terdeteksi dari data: Kolom {idx + 1} (berisi {len(value_str)} digit)")
                
                elif detected_columns['kpj'] == -1 and len(value_str) >= 10 and any(c.isalpha() for c in value_str):
                    detected_columns['kpj'] = idx
                    print(f"✅ Kolom KPJ terdeteksi dari data: Kolom {idx + 1}")
                
                elif detected_columns['nama'] == -1 and len(value_str) > 0 and any(c.isalpha() for c in value_str) and not value_str.isdigit():
                    detected_columns['nama'] = idx
                    print(f"✅ Kolom NAMA terdeteksi dari data: Kolom {idx + 1}")
    
    print("=" * 50)
    print("📊 HASIL DETEKSI KOLOM:")
    print(f"   NIK  : Kolom {detected_columns['nik'] + 1 if detected_columns['nik'] != -1 else 'TIDAK DITEMUKAN'}")
    print(f"   KPJ  : Kolom {detected_columns['kpj'] + 1 if detected_columns['kpj'] != -1 else 'TIDAK DITEMUKAN'}")
    print(f"   NAMA : Kolom {detected_columns['nama'] + 1 if detected_columns['nama'] != -1 else 'TIDAK DITEMUKAN'}")
    print("=" * 50)

def log(message, is_success=True):
    timestamp = datetime.datetime.now().strftime("%H:%M:%S")
    prefix = "✅" if is_success else "❌"
    print(f"{timestamp} {prefix} {message}")

from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.support.ui import WebDriverWait

def init_driver():
    global driver, wait

    opt = Options()
    
    # 🔥 WAJIB
    opt.binary_location = "/usr/bin/google-chrome-stable"

    opt.add_argument("--headless=new")
    opt.add_argument("--no-sandbox")
    opt.add_argument("--disable-dev-shm-usage")
    opt.add_argument("--disable-gpu")
    opt.add_argument("--window-size=1920,1080")
    
    # tambahan biar stabil
    opt.add_argument("--remote-debugging-port=9222")
    opt.add_argument("--disable-software-rasterizer")

    service = Service("/usr/bin/chromedriver")

    driver = webdriver.Chrome(service=service, options=opt)
    wait = WebDriverWait(driver, 15)

    return driver
    
def open_lasik_page():
    global driver
    print("🌐 Membuka halaman LASIK...")
    driver.get(LAPAKASIK_URL)
    time.sleep(5)
    
    close_selectors = [
        (By.ID, "btn-close-popup-banner"),
        (By.CLASS_NAME, "close"),
        (By.XPATH, "//button[contains(text(),'Tutup')]")
    ]
    
    for by, selector in close_selectors:
        try:
            btn = WebDriverWait(driver, 2).until(EC.element_to_be_clickable((by, selector)))
            driver.execute_script("arguments[0].click();", btn)
            time.sleep(0.5)
            break
        except:
            continue
    
    WebDriverWait(driver, 15).until(
        EC.presence_of_element_located((By.XPATH, "//input[@placeholder='Isi Nomor E-KTP']"))
    )
    print("✅ Halaman LASIK siap")

def get_field(placeholder):
    return wait.until(EC.element_to_be_clickable((By.XPATH, f"//input[@placeholder='{placeholder}']")))

def send_keys_with_retry(placeholder, text, attempts=3, wait_ms=300):
    if not text or text == 'nan':
        return
        
    for i in range(attempts):
        try:
            el = get_field(placeholder)
            driver.execute_script("arguments[0].value = '';", el)
            el.send_keys(str(text))
            return
        except (StaleElementReferenceException, WebDriverException):
            time.sleep(wait_ms / 1000)
    
    el = get_field(placeholder)
    driver.execute_script("arguments[0].value = '';", el)
    el.send_keys(str(text))

def click_with_retry(by, selector, attempts=3, wait_ms=200):
    for i in range(attempts):
        try:
            wait.until(EC.element_to_be_clickable((by, selector))).click()
            return
        except:
            time.sleep(wait_ms / 1000)
    
    driver.find_element(by, selector).click()

def clear_form_fields():
    try:
        fields = driver.find_elements(By.XPATH, "//input[@placeholder]")
        for field in fields:
            try:
                if field.is_displayed():
                    driver.execute_script("arguments[0].value = '';", field)
            except:
                pass
        time.sleep(0.5)
    except:
        pass

# ============================================================
# 🔥 CORE METHOD: EXTRACT SALDO JHT
# ============================================================
def extract_saldo_jht():
    """Extract saldo JHT"""
    try:
        print("Mencari saldo JHT...")
        
        try:
            WebDriverWait(driver, 15).until(
                EC.presence_of_element_located((By.XPATH, "//*[contains(text(),'JHT') or contains(text(),'Rp')]"))
            )
        except TimeoutException:
            print("Timeout menunggu elemen JHT/Rp")
            return "Rp 0"
        
        time.sleep(2)
        
        # STRATEGI 1: Cari elemen JHT dan Rp bersama
        jht_elements = driver.find_elements(By.XPATH, "//*[contains(text(),'JHT') and contains(text(),'Rp')]")
        jht_elements = [el for el in jht_elements if el.is_displayed() and el.text.strip()]
        
        for el in jht_elements:
            text = el.text.strip()
            match = re.search(r'Rp\s*([\d.,]+)', text)
            if match:
                angka_str = match.group(1).replace(".", "").replace(",", "")
                try:
                    angka = int(angka_str)
                    saldo_formatted = f"Rp {angka:,}"
                    print(f"✅ SALDO JHT DITEMUKAN: {saldo_formatted}")
                    return saldo_formatted
                except:
                    pass
        
        # STRATEGI 2: Cari semua elemen dengan Rp
        all_elements = driver.find_elements(By.XPATH, "//*[contains(text(),'Rp')]")
        all_elements = [el for el in all_elements if el.is_displayed() and el.text.strip()]
        
        semua_saldo = []
        
        for el in all_elements:
            text = el.text.strip()
            match = re.search(r'Rp\s*([\d.,]+)', text)
            if match:
                angka_str = match.group(1).replace(".", "").replace(",", "")
                try:
                    angka = int(angka_str)
                    semua_saldo.append(angka)
                except:
                    pass
        
        if semua_saldo:
            jht = max(semua_saldo)
            saldo_formatted = f"Rp {jht:,}"
            print(f"✅ SALDO JHT (terbesar): {saldo_formatted}")
            return saldo_formatted
        
        # STRATEGI 3: JavaScript injection
        print("🔄 Mencoba JavaScript injection...")
        
        js_result = driver.execute_script("""
            var elements = document.querySelectorAll('*');
            for(var i=0; i<elements.length; i++) {
                if(elements[i].__vue__) {
                    var vm = elements[i].__vue__;
                    if(vm.saldoJHT) return 'Rp ' + vm.saldoJHT;
                    if(vm.totalSaldo) return 'Rp ' + vm.totalSaldo;
                }
            }
            
            var bodyText = document.body.innerText;
            var match = bodyText.match(/JHT.*?Rp\\s*([\\d.,]+)/i);
            if(match) return 'Rp ' + match[1];
            
            return null;
        """)
        
        if js_result and js_result != "null":
            print(f"✅ SALDO JHT dari JS: {js_result}")
            return js_result
        
        print("⚠️ SALDO JHT TIDAK DITEMUKAN")
        return "Rp 0"
        
    except TimeoutException:
        print("⚠️ Timeout menunggu response")
        return "Rp 0"
    except Exception as ex:
        print(f"⚠️ Error extract saldo: {str(ex)}")
        return "Rp 0"

# ============================================================
# 🔥 POPUP DETECTION METHODS
# ============================================================
def is_konfirmasi_persetujuan_popup():
    try:
        popups = driver.find_elements(By.CSS_SELECTOR, ".swal2-popup.swal2-show")
        for p in popups:
            if not p.is_displayed():
                continue
            
            titles = p.find_elements(By.CSS_SELECTOR, ".swal2-title")
            if titles:
                title_text = titles[0].text.strip()
                if title_text == "Konfirmasi Persetujuan":
                    buttons = p.find_elements(By.CSS_SELECTOR, "button.swal2-confirm, button.swal2-cancel")
                    if len(buttons) >= 2:
                        return True
            
            contents = p.find_elements(By.CSS_SELECTOR, "#swal2-content")
            if contents:
                content = contents[0].text
                if "Tidak Bersedia" in content or "Bersedia" in content:
                    return True
        return False
    except:
        return False

def is_informasi_penting_popup():
    try:
        popups = driver.find_elements(By.CSS_SELECTOR, ".swal2-popup.swal2-show")
        for p in popups:
            if not p.is_displayed():
                continue
            titles = p.find_elements(By.CSS_SELECTOR, ".swal2-title")
            if titles and "Informasi Penting" in titles[0].text:
                return True
            contents = p.find_elements(By.CSS_SELECTOR, "#swal2-content")
            if contents and "Konfirmasi pengajuan klaim" in contents[0].text:
                return True
        return False
    except:
        return False

def is_pengaduan_klaim_popup():
    try:
        popups = driver.find_elements(By.CSS_SELECTOR, ".swal2-popup.swal2-show")
        for p in popups:
            if not p.is_displayed():
                continue
            contents = p.find_elements(By.CSS_SELECTOR, "#swal2-content")
            if contents:
                content_text = contents[0].text.lower()
                if "pengajuan klaim belum dapat dilanjutkan" in content_text or \
                   "hubungi petugas kantor cabang" in content_text:
                    return True
        return False
    except:
        return False

def is_jmo_redirect_popup():
    try:
        popups = driver.find_elements(By.CSS_SELECTOR, ".swal2-popup.swal2-show")
        for p in popups:
            if not p.is_displayed():
                continue
            content = p.text.lower()
            if "konfirmasi persetujuan" in content:
                continue
            if "jmo" in content or "jamsostek mobile" in content or "aplikasi jmo" in content:
                return True
        return False
    except:
        return False

def is_non_consent_popup():
    try:
        popups = driver.find_elements(By.CSS_SELECTOR, ".swal2-popup.swal2-show")
        for p in popups:
            if not p.is_displayed():
                continue
            
            content = p.text.lower()
            title = ""
            titles = p.find_elements(By.CSS_SELECTOR, ".swal2-title")
            if titles:
                title = titles[0].text.lower()
            
            if "konfirmasi persetujuan" in title or "konfirmasi persetujuan" in content:
                continue
            if "jmo" in content or "jamsostek mobile" in content:
                continue
            if "informasi penting" in title or "informasi penting" in content:
                continue
            if "pengajuan klaim" in content or "kantor cabang" in content:
                continue
            
            error_icons = p.find_elements(By.CSS_SELECTOR, ".swal2-icon.swal2-error")
            if error_icons and error_icons[0].is_displayed():
                return True
            
            contents = p.find_elements(By.CSS_SELECTOR, "#swal2-content")
            if contents:
                content_text = contents[0].text.lower()
                keywords = ["nomor", "sesuai", "kartu kepesertaan", "harap masukkan", 
                           "tidak valid", "salah", "gagal", "error", "maaf", 
                           "tidak ditemukan", "invalid", "format"]
                for keyword in keywords:
                    if keyword in content_text:
                        return True
            return True
        return False
    except:
        return False

def click_empty_area_to_close_popup():
    try:
        containers = driver.find_elements(By.CSS_SELECTOR, ".swal2-container")
        for container in containers:
            if container.is_displayed():
                driver.execute_script("""
                    var container = arguments[0];
                    var rect = container.getBoundingClientRect();
                    var clickEvent = new MouseEvent('click', {
                        view: window,
                        bubbles: true,
                        cancelable: true,
                        clientX: rect.left + 10,
                        clientY: rect.top + 10
                    });
                    container.dispatchEvent(clickEvent);
                """, container)
                time.sleep(0.1)
                return
        driver.execute_script("""
            var clickEvent = new MouseEvent('click', {
                view: window,
                bubbles: true,
                cancelable: true,
                clientX: 10,
                clientY: 10
            });
            document.body.dispatchEvent(clickEvent);
        """)
        time.sleep(0.1)
    except Exception as ex:
        print(f"⚠️ Error klik area kosong: {str(ex)}")

def handle_consent_popup():
    """Handle popup Konfirmasi Persetujuan"""
    try:
        popup = driver.find_element(By.CSS_SELECTOR, ".swal2-popup.swal2-show")
        tidak_bersedia = popup.find_elements(By.XPATH, ".//button[contains(text(),'Tidak Bersedia')]")
        
        if tidak_bersedia:
            driver.execute_script("arguments[0].click();", tidak_bersedia[0])
            time.sleep(2)
            return extract_saldo_jht()
        
        return "Rp 0"
    except Exception as ex:
        print(f"Error handle consent: {str(ex)}")
        return "Rp 0"

def handle_informasi_penting_popup(nik, nama):
    try:
        popup = driver.find_element(By.CSS_SELECTOR, ".swal2-popup.swal2-show")
        title = popup.find_element(By.CSS_SELECTOR, ".swal2-title").text
        content = popup.find_element(By.CSS_SELECTOR, "#swal2-content").text
        
        print("\n" + "="*60)
        print(f"⚠️ POPUP INFORMASI PENTING")
        print(f"👤 NIK: {nik}")
        print(f"👤 Nama: {nama}")
        print(f"📋 Title: {title}")
        print(f"📋 Pesan: {content.replace(chr(10), ' ').strip()}")
        print("="*60 + "\n")
        
        click_empty_area_to_close_popup()
    except Exception as ex:
        print(f"⚠️ Error handle popup informasi: {str(ex)}")

def handle_pengaduan_klaim_popup(nik, nama):
    try:
        popup = driver.find_element(By.CSS_SELECTOR, ".swal2-popup.swal2-show")
        content = popup.find_element(By.CSS_SELECTOR, "#swal2-content").text
        
        print("\n" + "="*60)
        print(f"⚠️ POPUP PENGADUAN KLAIM")
        print(f"👤 NIK: {nik}")
        print(f"👤 Nama: {nama}")
        print(f"📋 Pesan: {content}")
        print("="*60 + "\n")
        
        click_empty_area_to_close_popup()
    except Exception as ex:
        print(f"⚠️ Error handle popup pengaduan: {str(ex)}")

def handle_jmo_redirect_popup(nik, nama):
    try:
        popup = driver.find_element(By.CSS_SELECTOR, ".swal2-popup.swal2-show")
        content = popup.text
        
        print("\n" + "="*60)
        print(f"📱 POPUP ARAHAN JMO")
        print(f"👤 NIK: {nik}")
        print(f"👤 Nama: {nama}")
        print(f"📋 Pesan: {content}")
        print("="*60 + "\n")
        
        click_empty_area_to_close_popup()
    except Exception as ex:
        print(f"⚠️ Error handle popup arahan JMO: {str(ex)}")

def handle_non_consent_popup(nik, nama):
    try:
        popup = driver.find_element(By.CSS_SELECTOR, ".swal2-popup.swal2-show")
        title = popup.find_elements(By.CSS_SELECTOR, ".swal2-title")
        title_text = title[0].text if title else "Tidak ada title"
        content = popup.find_element(By.CSS_SELECTOR, "#swal2-content").text
        
        print("\n" + "="*60)
        print(f"⚠️ POPUP NON-CONSENT")
        print(f"👤 NIK: {nik}")
        print(f"👤 Nama: {nama}")
        print(f"📋 Title: {title_text}")
        print(f"📋 Pesan: {content}")
        print("="*60 + "\n")
        
        click_empty_area_to_close_popup()
    except Exception as ex:
        print(f"⚠️ Error handle popup non-consent: {str(ex)}")

def handle_nomor_kartu_popup(nik, nama):
    try:
        popup = driver.find_element(By.CSS_SELECTOR, ".swal2-popup.swal2-show")
        content = popup.find_element(By.CSS_SELECTOR, "#swal2-content").text
        
        print("\n" + "="*60)
        print(f"⚠️ POPUP VALIDASI NOMOR KARTU")
        print(f"👤 NIK: {nik}")
        print(f"👤 Nama: {nama}")
        print(f"📋 Pesan: {content}")
        print("="*60 + "\n")
        
        click_empty_area_to_close_popup()
    except Exception as ex:
        print(f"⚠️ Error handle popup nomor kartu: {str(ex)}")

def cleanup_all_popups():
    try:
        popups = driver.find_elements(By.CSS_SELECTOR, ".swal2-popup.swal2-show")
        if popups:
            click_empty_area_to_close_popup()
    except:
        pass

def extract_nama_pt():
    try:
        print("🔍 Mencari Nama Perusahaan (PT)...")
        time.sleep(1)
        
        pt_elements = driver.find_elements(By.XPATH,
            "//*[contains(text(), 'PT ') or contains(text(), 'CV ') or contains(text(), 'Perusahaan') or contains(text(), 'PERSERO')]")
        pt_elements = [el for el in pt_elements if el.is_displayed() and el.text.strip()]
        
        for el in pt_elements:
            text = el.text.strip()
            if 5 < len(text) < 100:
                text = re.sub(r'\s+', ' ', text)
                if text.startswith("PT ") or text.startswith("PT."):
                    print(f"📋 Menemukan: {text}")
                    return text
        
        perusahaan_elements = driver.find_elements(By.XPATH,
            "//td[contains(text(), 'Perusahaan')]/following-sibling::td | " +
            "//th[contains(text(), 'Perusahaan')]/following-sibling::td | " +
            "//div[contains(@class, 'perusahaan')]")
        
        for el in perusahaan_elements:
            text = el.text.strip()
            if text and len(text) > 5:
                print(f"📋 Menemukan (dari tabel): {text}")
                return text
        
        all_texts = driver.find_elements(By.XPATH, "//*[contains(text(), 'PT')]")
        all_texts = [el.text.strip() for el in all_texts if el.is_displayed() and len(el.text) > 10]
        
        for text in all_texts:
            if "PT " in text and len(text) < 100:
                print(f"📋 Menemukan (alternatif): {text}")
                return text
        
        print("⚠️ Nama Perusahaan tidak ditemukan")
        return "-"
    except Exception as ex:
        print(f"⚠️ Error ExtractNamaPT: {str(ex)}")
        return "-"

def detect_jmo_status(nik, nama, kelurahan):
    try:
        if nik and "MIG" in str(nik).upper():
            return "NIK MIGRASI"
        if not nik or str(nik) == "-" or str(nik) == "nan":
            return "NIK KOSONG"
        if kelurahan and kelurahan != "nan" and ("TIDAK TERDAFTAR" in str(kelurahan).upper() or "BELUM TERDAFTAR" in str(kelurahan).upper()):
            return "TIDAK ADA DPT"
        return "CEK POPUP..."
    except Exception as ex:
        print(f"⚠️ Error deteksi JMO: {str(ex)}")
        return "ERROR DETEKSI"

# ============================================================
# 🔥 FUNGSI SHOW CONFIRMATION - TIDAK MENGHENTIKAN PROSES
# ============================================================
def show_console_confirmation(message):
    """Tampilkan konfirmasi di console - TANPA MENUNGGU INPUT"""
    global AUTO_MODE
    
    try:
        print("\n" + "█"*80)
        print("██" + " "*76 + "██")
        
        lines = message.split('\n')
        for line in lines:
            centered_line = line.ljust(76)[:76]
            print(f"██ {centered_line} ██")
        
        print("██" + " "*76 + "██")
        
        if AUTO_MODE:
            print("██ " + "🔄 AUTO MODE: LANJUT KE DATA BERIKUTNYA...".ljust(76) + " ██")
        else:
            print("██ " + "TEKAN ENTER UNTUK MELANJUTKAN...".ljust(76) + " ██")
        
        print("██" + " "*76 + "██")
        print("█"*80 + "\n")
        
        # 🔥 JANGAN PAKAI input() - LANGSUNG LANJUT
        if not AUTO_MODE:
            try:
                input()
            except:
                pass
        else:
            time.sleep(1)  # Jeda singkat agar terbaca
        
        return True
    except Exception as ex:
        print(f"⚠️ Error: {str(ex)}")
        return False

def process_row(row_num, kpj, nama, nik, kelurahan):
    """Proses satu baris data"""
    global sukses_count, gagal_count
    
    try:
        status_jmo = detect_jmo_status(nik, nama, kelurahan)
        
        if status_jmo in ["NIK MIGRASI", "NIK KOSONG", "TIDAK ADA DPT"]:
            print(f"⏭️ [SKIP] {nik} - {status_jmo}")
            return False, "SKIP", "Rp 0", status_jmo, "-"
        
        # Input data
        if nik and str(nik) != "nan":
            send_keys_with_retry("Isi Nomor E-KTP", nik)
        if kpj and str(kpj) != "nan":
            send_keys_with_retry("Isi Nomor KPJ", kpj)
        if nama and str(nama) != "nan":
            send_keys_with_retry("Isi Nama sesuai KTP", nama)
        
        click_with_retry(By.TAG_NAME, "body")
        time.sleep(0.9)
        
        # Submit
        try:
            submit_buttons = driver.find_elements(By.XPATH,
                "//button[contains(text(),'Cari')] | " +
                "//button[contains(text(),'Submit')] | " +
                "//button[@type='submit']")
            
            if submit_buttons:
                for btn in submit_buttons:
                    if btn.is_displayed() and btn.is_enabled():
                        driver.execute_script("arguments[0].click();", btn)
                        break
            else:
                driver.find_element(By.TAG_NAME, "body").send_keys("\n")
        except:
            pass
        
        time.sleep(2.5)
        
        # Cek popup Konfirmasi Persetujuan
        if is_konfirmasi_persetujuan_popup():
            saldo = handle_consent_popup()
            nama_pt = extract_nama_pt()
            return True, "SUKSES", saldo, "AKTIF", nama_pt
        
        # Cek popup Informasi Penting
        if is_informasi_penting_popup():
            handle_informasi_penting_popup(nik, nama)
            return False, "SKIP", "Rp 0", "SKIP - INFORMASI PENTING", "-"
        
        # Cek popup Pengaduan Klaim
        if is_pengaduan_klaim_popup():
            handle_pengaduan_klaim_popup(nik, nama)
            return False, "GAGAL", "Rp 0", "BLOKIR KLAIM", "-"
        
        # Cek popup arahan JMO
        if is_jmo_redirect_popup():
            handle_jmo_redirect_popup(nik, nama)
            return False, "GAGAL", "Rp 0", "BELUM REGISTRASI JMO", "-"
        
        # Cek popup non-consent
        if is_non_consent_popup():
            try:
                popup = driver.find_element(By.CSS_SELECTOR, ".swal2-popup.swal2-show")
                content = popup.find_element(By.CSS_SELECTOR, "#swal2-content").text.lower()
                
                if "nomor" in content and "kartu kepesertaan" in content:
                    handle_nomor_kartu_popup(nik, nama)
                else:
                    handle_non_consent_popup(nik, nama)
            except:
                handle_non_consent_popup(nik, nama)
            
            return False, "GAGAL", "Rp 0", "GAGAL - POPUP VALIDASI", "-"
        
        # Tidak ada popup
        return False, "GAGAL", "Rp 0", "TIDAK ADA RESPON", "-"
        
    except Exception as ex:
        print(f"⚠️ {nik}: Error process - {str(ex)}")
        return False, "ERROR", "Rp 0", "ERROR SISTEM", "-"
    finally:
        cleanup_all_popups()

def save_to_excel(output_path, excel_file_path):
    global processed_rows
    
    try:
        wb = load_workbook(excel_file_path)
        ws = wb.active
        
        from openpyxl import Workbook
        output_wb = Workbook()
        output_ws = output_wb.active
        output_ws.title = "HASIL"
        
        for col in range(1, ws.max_column + 1):
            output_ws.cell(row=1, column=col, value=ws.cell(row=1, column=col).value)
        
        output_ws.cell(row=1, column=ws.max_column + 1, value="SALDO JHT")
        output_ws.cell(row=1, column=ws.max_column + 2, value="NAMA PERUSAHAAN")
        
        data_dengan_saldo = [row for row in processed_rows if row.saldo_jht != "Rp 0"]
        
        print(f"📊 Menyimpan {len(data_dengan_saldo)} data dengan SALDO (dari total {len(processed_rows)})")
        
        output_row = 2
        for row_data in data_dengan_saldo:
            for col in range(1, ws.max_column + 1):
                if col in row_data.all_columns_data:
                    output_ws.cell(row=output_row, column=col, value=row_data.all_columns_data[col])
                else:
                    output_ws.cell(row=output_row, column=col, value=ws.cell(row=row_data.row_number, column=col).value)
            
            output_ws.cell(row=output_row, column=ws.max_column + 1, value=row_data.saldo_jht)
            output_ws.cell(row=output_row, column=ws.max_column + 2, value=row_data.nama_pt)
            output_row += 1
        
        output_wb.save(output_path)
        
        print("\n" + "="*70)
        print("📊 RINGKASAN HASIL:")
        print("-"*70)
        print(f"✅ Data dengan SALDO: {len(data_dengan_saldo)} data (DISIMPAN)")
        print(f"❌ Data tanpa saldo: {len(processed_rows) - len(data_dengan_saldo)} data (TIDAK DISIMPAN)")
        print(f"📊 TOTAL DIPROSES: {len(processed_rows)} data")
        
        if data_dengan_saldo:
            print("\n📋 DETAIL DATA YANG DISIMPAN:")
            for row in data_dengan_saldo:
                print(f"✅ {row.nik} - {row.nama}: {row.saldo_jht}")
        
        print("="*70 + "\n")
        
        print(f"💾 Excel berhasil disimpan: {output_path}")
        
    except Exception as ex:
        print(f"❌ Error save Excel: {str(ex)}")

# ===============================
# FLASK ROUTES
# ===============================

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/upload', methods=['POST'])
def upload_file():
    global excel_path, processed_rows, sukses_count, gagal_count
    
    processed_rows = []
    sukses_count = 0
    gagal_count = 0
    
    if 'file' not in request.files:
        return jsonify({"error": "No file"}), 400
    
    file = request.files['file']
    if file.filename == '':
        return jsonify({"error": "Empty file"}), 400
    
    filepath = os.path.join(UPLOAD_FOLDER, file.filename)
    file.save(filepath)
    excel_path = filepath
    
    try:
        df = pd.read_excel(filepath)
        detect_columns(df)
        
        if detected_columns['nik'] == -1 and detected_columns['kpj'] == -1:
            return jsonify({"error": "Tidak dapat mendeteksi kolom NIK atau KPJ!"}), 400
        
        total = len(df)
        return jsonify({
            "total": total,
            "detected_columns": {
                "nik": detected_columns['nik'] + 1 if detected_columns['nik'] != -1 else None,
                "kpj": detected_columns['kpj'] + 1 if detected_columns['kpj'] != -1 else None,
                "nama": detected_columns['nama'] + 1 if detected_columns['nama'] != -1 else None
            },
            "message": f"Loaded {total} records"
        })
    except Exception as e:
        return jsonify({"error": str(e)}), 400

@app.route('/start')
def start():
    def generate():
        global driver, is_running, stop_flag, sukses_count, gagal_count, processed_rows, current_output_path, excel_path, detected_columns
        
        is_running = True
        stop_flag = False
        
        try:
            yield f"data: {json.dumps({'type': 'log', 'message': 'Initializing Chrome driver...'})}\n\n"
            driver = init_driver()
            yield f"data: {json.dumps({'type': 'log', 'message': 'Chrome driver ready'})}\n\n"
            
            yield f"data: {json.dumps({'type': 'log', 'message': 'Opening LASIK page...'})}\n\n"
            open_lasik_page()
            yield f"data: {json.dumps({'type': 'log', 'message': 'LASIK page ready'})}\n\n"
            
        except Exception as e:
            yield f"data: {json.dumps({'type': 'error', 'message': f'Chrome error: {str(e)}'})}\n\n"
            return
        
        if not excel_path or not os.path.exists(excel_path):
            yield f"data: {json.dumps({'type': 'error', 'message': 'No file uploaded'})}\n\n"
            return
        
        try:
            df = pd.read_excel(excel_path)
            total_rows = len(df) + 1
            
            current_output_path = os.path.join(RESULT_FOLDER, f"HASIL_JMO_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
            
            yield f"data: {json.dumps({'type': 'start', 'total': total_rows - 1})}\n\n"
            
            print("\n" + "="*70)
            print("🚀 BOT LASIK - DETEKSI STATUS JMO")
            print("✅ AUTO MODE: LANGSUNG PROSES TANPA KONFIRMASI")
            print("="*70 + "\n")
            
            for idx, row in df.iterrows():
                if stop_flag or not is_running:
                    yield f"data: {json.dumps({'type': 'stop', 'message': 'Process stopped'})}\n\n"
                    break
                
                row_num = idx + 2
                
                nik = ""
                kpj = ""
                nama = ""
                kelurahan = ""
                
                if detected_columns['nik'] != -1:
                    nik = str(row.iloc[detected_columns['nik']]).strip() if pd.notna(row.iloc[detected_columns['nik']]) else ""
                
                if detected_columns['kpj'] != -1:
                    kpj = str(row.iloc[detected_columns['kpj']]).strip() if pd.notna(row.iloc[detected_columns['kpj']]) else ""
                
                if detected_columns['nama'] != -1:
                    nama = str(row.iloc[detected_columns['nama']]).strip() if pd.notna(row.iloc[detected_columns['nama']]) else ""
                
                if detected_columns['kelurahan'] != -1:
                    kelurahan = str(row.iloc[detected_columns['kelurahan']]).strip() if pd.notna(row.iloc[detected_columns['kelurahan']]) else ""
                
                # FALLBACK: cari NIK di semua kolom
                if (not nik or nik == "nan") and detected_columns['nik'] == -1:
                    for col_idx, val in enumerate(row):
                        val_str = str(val).strip() if pd.notna(val) else ""
                        if len(val_str) >= 16 and val_str.isdigit():
                            nik = val_str
                            print(f"🔍 NIK ditemukan di kolom {col_idx + 1}: {nik}")
                            break
                
                if (not nik or nik == "nan") and (not kpj or kpj == "nan"):
                    gagal_count += 1
                    print(f"⏭️ Skip baris {row_num}: NIK dan KPJ kosong")
                    continue
                
                progress = idx + 1
                percent = (progress / (total_rows - 1)) * 100
                yield f"data: {json.dumps({'type': 'progress', 'current': progress, 'total': total_rows - 1, 'percent': percent, 'kpj': kpj})}\n\n"
                
                print(f"🔄 Processing {progress}/{total_rows - 1} - NIK: {nik}, KPJ: {kpj}")
                
                all_columns_data = {}
                for col_idx in range(len(df.columns)):
                    all_columns_data[col_idx + 1] = row.iloc[col_idx] if pd.notna(row.iloc[col_idx]) else ""
                
                try:
                    success, status, saldo, status_jmo, nama_pt = process_row(row_num, kpj, nama, nik, kelurahan)
                    
                    processed_rows.append(ProcessedRow(row_num, all_columns_data, kpj, nama, nik, status, saldo, status_jmo, nama_pt))
                    
                    if success:
                        sukses_count += 1
                        print(f"✅ {nik} | {status_jmo} | {saldo}")
                    else:
                        gagal_count += 1
                        print(f"❌ {nik} - {status_jmo}")
                    
                    # Tampilkan ringkasan singkat di console (tanpa menunggu)
                    print(f"📊 Hasil: {nik} - {status_jmo} - {saldo}")
                    
                    clear_form_fields()
                    time.sleep(0.9)
                    
                    if (sukses_count + gagal_count) % 5 == 0 and (sukses_count + gagal_count) > 0:
                        save_to_excel(current_output_path, excel_path)
                        yield f"data: {json.dumps({'type': 'log', 'message': f'Auto-saved at {sukses_count + gagal_count} records'})}\n\n"
                    
                    result_data = {
                        'kpj': kpj,
                        'nama': nama,
                        'nik': nik,
                        'status': status,
                        'keterangan': status_jmo,
                        'saldo_jht': saldo,
                        'nama_perusahaan': nama_pt
                    }
                    yield f"data: {json.dumps({'type': 'result', 'data': result_data})}\n\n"
                    
                except Exception as ex:
                    gagal_count += 1
                    print(f"⚠️ Error: {nik} - {str(ex)}")
                    yield f"data: {json.dumps({'type': 'error', 'message': f'Error: {str(ex)[:100]}'})}\n\n"
            
            if processed_rows:
                save_to_excel(current_output_path, excel_path)
                
                print("\n" + "="*70)
                print(f"🎯 FINISH | Success: {sukses_count} | Failed: {gagal_count}")
                print(f"💰 Total data dengan saldo: {len([r for r in processed_rows if r.saldo_jht != 'Rp 0'])}")
                print("="*70 + "\n")
                
                yield f"data: {json.dumps({'type': 'done', 'total': len(processed_rows), 'message': f'Completed! {sukses_count} success, {gagal_count} failed'})}\n\n"
            else:
                yield f"data: {json.dumps({'type': 'done', 'total': 0, 'message': 'No records processed'})}\n\n"
                
        except Exception as e:
            yield f"data: {json.dumps({'type': 'error', 'message': f'System error: {str(e)}'})}\n\n"
        finally:
            is_running = False
            try:
                if driver:
                    driver.quit()
            except:
                pass
    
    return Response(generate(), mimetype='text/event-stream')

@app.route('/stop', methods=['POST'])
def stop():
    global stop_flag, is_running
    stop_flag = True
    is_running = False
    return jsonify({"status": "stopped"})

@app.route('/download')
def download():
    global current_output_path
    if not current_output_path or not os.path.exists(current_output_path):
        return jsonify({"error": "No result file available"}), 404
    return send_file(current_output_path, as_attachment=True, download_name=os.path.basename(current_output_path))

@app.route('/token-status')
def token_status():
    return jsonify({"active": True})

@app.route('/get-balance')
def get_balance():
    return jsonify({"balance": "N/A"})

@app.route('/open-oss', methods=['POST'])
def open_oss():
    return jsonify({"msg": "LASIK mode active"})

if __name__ == '__main__':
    print("\n" + "="*60)
    print("🚀 BPJS LASIK BOT - AUTO MODE")
    print("="*60)
    print("📋 URL: http://localhost:5000")
    print("📁 Results folder: results/")
    print("="*60)
    print("\n✨ PERBAIKAN:")
    print("   1. ✅ AUTO MODE - LANGSUNG PROSES TANPA KONFIRMASI")
    print("   2. ✅ TIDAK PERLU TEKAN ENTER LAGI")
    print("   3. ✅ PROSES LANJUT OTOMATIS KE DATA BERIKUTNYA")
    print("   4. ✅ DETEKSI KOLOM OTOMATIS")
    print("   5. ✅ EKSTRAK SALDO JHT (3 STRATEGI)")
    print("="*60)
    print("\n📊 FORMAT EXCEL FLEKSIBEL:")
    print("   - NIK: cari kata kunci 'nik', 'ktp' ATAU angka 16 digit")
    print("   - KPJ: cari kata kunci 'kpj', 'peserta' ATAU kombinasi angka+huruf")
    print("   - NAMA: cari kata kunci 'nama', 'name' ATAU teks non-angka")
    print("="*60 + "\n")
    
    app.run(debug=False, host='0.0.0.0', port=5000, threaded=True)
