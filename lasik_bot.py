import os
import re
import time
import subprocess
from datetime import datetime
from typing import Optional, Tuple, List, Dict, Callable
from openpyxl import load_workbook
from openpyxl import Workbook
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import (
    StaleElementReferenceException,
    WebDriverException,
    TimeoutException,
    NoSuchElementException
)


class LasikBotService:
    def __init__(self):
        self.excel_path: Optional[str] = None
        self.driver: Optional[webdriver.Chrome] = None
        self.wait: Optional[WebDriverWait] = None
        self.is_running: bool = False
        self._on_progress: Optional[Callable] = None
        self._on_status: Optional[Callable] = None
        self.sukses_count: int = 0
        self.gagal_count: int = 0
        self.current_output_path: Optional[str] = None
        self.processed_rows: List[Dict] = []
        self.LAPAKASIK_URL = "https://jmoantrian.bpjsketenagakerjaan.go.id/?source=e419a6aed6c50fefd9182774c25450b333de8d5e29169de6018bd1abb1c8f89b"

    @property
    def success_count(self) -> int:
        return self.sukses_count

    @property
    def failed_count(self) -> int:
        return self.gagal_count

    def set_on_progress(self, callback: Callable[[int, int], None]):
        """Set callback untuk progress update"""
        self._on_progress = callback

    def set_on_status(self, callback: Callable[[str, bool], None]):
        """Set callback untuk status update"""
        self._on_status = callback

    def _log(self, message: str, level: str = "info"):
        """Logging dengan level"""
        is_success = level != "error"
        if self._on_status:
            self._on_status(message, is_success)

        prefix = {
            "success": "[SUCCESS]",
            "warning": "[WARNING]",
            "error": "[ERROR]",
            "info": "[INFO]"
        }.get(level, "[INFO]")

        print(f"{datetime.now().strftime('%H:%M:%S')} {prefix} {message}")

    def open_chrome_only(self):
        """Membuka Chrome dengan debug port"""
        chrome_path = r"C:\Program Files\Google\Chrome\Application\chrome.exe"
        
        # Cek apakah Chrome terinstall di path default
        if not os.path.exists(chrome_path):
            # Coba cari di Program Files (x86)
            chrome_path = r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe"
            if not os.path.exists(chrome_path):
                self._log("Chrome tidak ditemukan di path default", "error")
                raise Exception("Chrome tidak ditemukan")
        
        args = [
            chrome_path,
            "--remote-debugging-port=9222",
            f"--user-data-dir=C:\\ChromeProfileLasik",
            self.LAPAKASIK_URL
        ]
        
        try:
            subprocess.Popen(args)
            time.sleep(3)
            self._log("Chrome berhasil dibuka", "success")
        except Exception as e:
            self._log(f"Gagal membuka Chrome: {str(e)}", "error")
            raise

    def attach_bot(self):
        """Attach ke Chrome yang sudah berjalan"""
        try:
            options = webdriver.ChromeOptions()
            options.debugger_address = "127.0.0.1:9222"
            
            # Tambahkan options untuk menghindari deteksi
            options.add_experimental_option("excludeSwitches", ["enable-logging"])
            options.add_argument("--disable-blink-features=AutomationControlled")
            options.add_argument("--disable-extensions")
            options.add_argument("--no-sandbox")
            options.add_argument("--disable-dev-shm-usage")
            
            self.driver = webdriver.Chrome(options=options)
            self.driver.implicitly_wait(0)
            self.wait = WebDriverWait(self.driver, 15)
            
            self._log("Bot berhasil terhubung ke Chrome", "success")
            if self._on_status:
                self._on_status("🤖 Bot berhasil terhubung ke Chrome", True)
        except Exception as ex:
            error_msg = f"Gagal attach Chrome: {str(ex)}"
            self._log(error_msg, "error")
            if self._on_status:
                self._on_status(f"❌ {error_msg}", False)
            raise

    def _get_field(self, placeholder: str):
        """Mendapatkan element input berdasarkan placeholder"""
        return self.wait.until(
            EC.visibility_of_element_located(
                (By.XPATH, f"//input[@placeholder='{placeholder}']")
            )
        )

    def _send_keys_with_retry(self, placeholder: str, text: str, attempts: int = 3, wait_ms: int = 300):
        """SendKeys dengan retry mechanism"""
        for i in range(attempts):
            try:
                el = self._get_field(placeholder)
                el.clear()
                el.send_keys(text)
                return
            except (StaleElementReferenceException, WebDriverException):
                time.sleep(wait_ms / 1000)
        
        last = self._get_field(placeholder)
        last.clear()
        last.send_keys(text)

    def _click_with_retry(self, by: By, value: str, attempts: int = 3, wait_ms: int = 200):
        """Click dengan retry mechanism"""
        for i in range(attempts):
            try:
                element = self.wait.until(EC.element_to_be_clickable((by, value)))
                element.click()
                return
            except:
                time.sleep(wait_ms / 1000)
        
        self.driver.find_element(by, value).click()

    def _clear_form_fields(self):
        """Membersihkan semua form fields"""
        try:
            fields = self.driver.find_elements(By.XPATH, "//input[@placeholder]")
            for field in fields:
                try:
                    if field.is_displayed():
                        field.clear()
                except:
                    pass
            time.sleep(0.5)
        except:
            pass

    def _extract_saldo_jht(self) -> str:
        """Ekstrak saldo JHT dari halaman"""
        try:
            self._log("Mencari saldo JHT...")
            
            # Tunggu hingga halaman load
            try:
                self.wait.until(lambda d: len(d.find_elements(By.XPATH, "//*[contains(text(),'JHT') or contains(text(),'Rp')]")) > 0)
            except:
                pass
            
            time.sleep(2)
            
            # Cari elemen yang mengandung JHT dan Rp
            jht_elements = [el for el in self.driver.find_elements(
                By.XPATH, "//*[contains(text(),'JHT') and contains(text(),'Rp')]"
            ) if el.is_displayed() and el.text.strip()]
            
            for el in jht_elements:
                text = el.text.strip()
                match = re.search(r'Rp\s*([\d.,]+)', text)
                if match:
                    angka_str = match.group(1).replace(".", "").replace(",", "")
                    if angka_str.isdigit():
                        angka = int(angka_str)
                        saldo_formatted = f"Rp {angka:,}"
                        self._log(f"Saldo JHT ditemukan: {saldo_formatted}", "success")
                        if self._on_status:
                            self._on_status(f"✅ SALDO JHT DITEMUKAN: {saldo_formatted}", True)
                        return saldo_formatted
            
            # Cari semua elemen dengan Rp
            all_elements = [el for el in self.driver.find_elements(
                By.XPATH, "//*[contains(text(),'Rp')]"
            ) if el.is_displayed() and el.text.strip()]
            
            semua_saldo = []
            for el in all_elements:
                text = el.text.strip()
                match = re.search(r'Rp\s*([\d.,]+)', text)
                if match:
                    angka_str = match.group(1).replace(".", "").replace(",", "")
                    if angka_str.isdigit():
                        semua_saldo.append(int(angka_str))
            
            if semua_saldo:
                jht = max(semua_saldo)
                saldo_formatted = f"Rp {jht:,}"
                self._log(f"Saldo JHT (terbesar): {saldo_formatted}", "success")
                if self._on_status:
                    self._on_status(f"✅ SALDO JHT (terbesar): {saldo_formatted}", True)
                return saldo_formatted
            
            # Fallback: coba dengan JavaScript
            result = self.driver.execute_script("""
                var elements = document.querySelectorAll('*');
                for(var i=0; i<elements.length; i++) {
                    if(elements[i].__vue__) {
                        var vm = elements[i].__vue__;
                        if(vm.saldoJHT) return 'Rp ' + vm.saldoJHT;
                        if(vm.totalSaldo) return 'Rp ' + vm.totalSaldo;
                    }
                }
                var bodyText = document.body.innerText;
                var match = bodyText.match(/JHT.*?Rp\\s*([\\d.,]+)/i);
                if(match) return 'Rp ' + match[1];
                return null;
            """)
            
            if result:
                self._log(f"Saldo JHT (JS): {result}", "success")
                if self._on_status:
                    self._on_status(f"✅ SALDO JHT : {result}", True)
                return str(result)
            
            self._log("Saldo JHT tidak ditemukan", "warning")
            if self._on_status:
                self._on_status("⚠️ SALDO JHT TIDAK DITEMUKAN", False)
            return "Rp 0"
            
        except TimeoutException:
            self._log("Timeout mendapatkan saldo JHT", "warning")
            if self._on_status:
                self._on_status("⚠️ Timeout mendapatkan saldo JHT", False)
            return "Rp 0"
        except Exception as ex:
            self._log(f"Error mendapatkan saldo: {str(ex)}", "error")
            if self._on_status:
                self._on_status(f"⚠️ Error mendapatkan saldo: {str(ex)}", False)
            return "Rp 0"

    def _extract_nama_pt(self) -> str:
        """Ekstrak nama perusahaan (PT)"""
        try:
            self._log("Mencari Nama Perusahaan (PT)...")
            time.sleep(1)
            
            # Cari elemen yang mengandung PT, CV, Perusahaan, PERSERO
            pt_elements = [el for el in self.driver.find_elements(
                By.XPATH, "//*[contains(text(), 'PT ') or contains(text(), 'CV ') or contains(text(), 'Perusahaan') or contains(text(), 'PERSERO')]"
            ) if el.is_displayed() and el.text.strip()]
            
            for el in pt_elements:
                text = el.text.strip()
                if 5 < len(text) < 100:
                    text = re.sub(r'\s+', ' ', text)
                    self._log(f"Menemukan: {text}")
                    if text.startswith("PT ") or text.startswith("PT."):
                        return text
            
            return "-"
        except Exception as ex:
            self._log(f"Error ExtractNamaPT: {str(ex)}", "error")
            return "-"

    def _detect_jmo_status(self, nik: str, nama: str, kelurahan: str) -> str:
        """Deteksi status JMO"""
        try:
            if nik and "MIG" in nik.upper():
                if self._on_status:
                    self._on_status(f"📱 NIK MIGRASI terdeteksi untuk {nik}", True)
                return "NIK MIGRASI"
            
            if not nik or nik == "-":
                if self._on_status:
                    self._on_status(f"📱 NIK KOSONG untuk {nama}", True)
                return "NIK KOSONG"
            
            if kelurahan and ("TIDAK TERDAFTAR" in kelurahan.upper() or "BELUM TERDAFTAR" in kelurahan.upper()):
                if self._on_status:
                    self._on_status(f"📱 TIDAK ADA DPT untuk {nik}", True)
                return "TIDAK ADA DPT"
            
            return "CEK POPUP..."
        except Exception as ex:
            if self._on_status:
                self._on_status(f"⚠️ Error deteksi JMO: {str(ex)}", False)
            return "ERROR DETEKSI"

    def _is_konfirmasi_persetujuan_popup(self) -> bool:
        """Deteksi popup Konfirmasi Persetujuan"""
        try:
            popups = self.driver.find_elements(By.CSS_SELECTOR, ".swal2-popup.swal2-show")
            for p in popups:
                if not p.is_displayed():
                    continue
                
                titles = p.find_elements(By.CSS_SELECTOR, ".swal2-title")
                if titles:
                    title_text = titles[0].text.strip()
                    if title_text == "Konfirmasi Persetujuan":
                        buttons = p.find_elements(By.CSS_SELECTOR, "button.swal2-confirm, button.swal2-cancel")
                        if len(buttons) >= 2:
                            return True
            return False
        except:
            return False

    def _click_empty_area_to_close_popup(self):
        """Klik area kosong untuk menutup popup"""
        try:
            containers = self.driver.find_elements(By.CSS_SELECTOR, ".swal2-container")
            for container in containers:
                if container.is_displayed():
                    self.driver.execute_script("""
                        var container = arguments[0];
                        var rect = container.getBoundingClientRect();
                        var clickEvent = new MouseEvent('click', {
                            view: window,
                            bubbles: true,
                            cancelable: true,
                            clientX: rect.left + 10,
                            clientY: rect.top + 10
                        });
                        container.dispatchEvent(clickEvent);
                    """, container)
                    time.sleep(0.1)
                    return
            
            # Alternatif: klik body
            self.driver.execute_script("""
                var clickEvent = new MouseEvent('click', {
                    view: window,
                    bubbles: true,
                    cancelable: true,
                    clientX: 10,
                    clientY: 10
                });
                document.body.dispatchEvent(clickEvent);
            """)
            time.sleep(0.1)
        except Exception as ex:
            if self._on_status:
                self._on_status(f"⚠️ Error klik area kosong: {str(ex)}", False)

    def _handle_consent_popup(self, nik: str, nama: str) -> str:
        """Handle popup konfirmasi persetujuan"""
        try:
            popup = self.driver.find_element(By.CSS_SELECTOR, ".swal2-popup.swal2-show")
            tidak_bersedia = popup.find_elements(By.XPATH, ".//button[contains(text(),'Tidak Bersedia')]")
            
            if tidak_bersedia:
                self.driver.execute_script("arguments[0].click();", tidak_bersedia[0])
                time.sleep(2)
                return self._extract_saldo_jht()
            
            return "Rp 0"
        except Exception as ex:
            self._log(f"Error handle consent popup: {str(ex)}", "error")
            return "Rp 0"

    def _process_row(self, row: int, kpj: str, nama: str, nik: str, kelurahan: str) -> Tuple[bool, str, str, str, str]:
        """Proses satu baris data"""
        popup_action = ""
        status = "GAGAL"
        saldo_jht = "Rp 0"
        status_jmo = "TIDAK DIKETAHUI"
        
        try:
            status_jmo = self._detect_jmo_status(nik, nama, kelurahan)
            
            if status_jmo in ["NIK MIGRASI", "NIK KOSONG", "TIDAK ADA DPT"]:
                if self._on_status:
                    self._on_status(f"⏭️ [SKIP] {nik} - {status_jmo}", True)
                return False, popup_action, status, saldo_jht, status_jmo
            
            # Input data
            self._send_keys_with_retry("Isi Nomor E-KTP", nik)
            self._send_keys_with_retry("Isi Nomor KPJ", kpj)
            self._send_keys_with_retry("Isi Nama sesuai KTP", nama)
            
            self._click_with_retry(By.TAG_NAME, "body")
            time.sleep(0.9)
            
            # Submit
            try:
                submit_buttons = self.driver.find_elements(By.XPATH, 
                    "//button[contains(text(),'Cari')] | //button[contains(text(),'Submit')] | //button[@type='submit']")
                
                if submit_buttons:
                    for btn in submit_buttons:
                        if btn.is_displayed() and btn.is_enabled():
                            self.driver.execute_script("arguments[0].click();", btn)
                            break
                else:
                    self.driver.find_element(By.TAG_NAME, "body").send_keys(Keys.ENTER)
            except:
                pass
            
            time.sleep(2.5)
            
            # Cek berbagai jenis popup
            if self._is_konfirmasi_persetujuan_popup():
                status = "SUKSES"
                popup_action = "KONFIRMASI_PERSETUJUAN"
                saldo_jht = self._handle_consent_popup(nik, nama)
                return True, popup_action, status, saldo_jht, status_jmo
            
            # Jika tidak ada popup sukses
            status = "GAGAL"
            popup_action = "NO_POPUP"
            status_jmo = "TIDAK ADA RESPON"
            return False, popup_action, status, saldo_jht, status_jmo
            
        except Exception as ex:
            status = "ERROR"
            popup_action = "EXCEPTION"
            status_jmo = "ERROR SISTEM"
            self._log(f"Error process row: {str(ex)}", "error")
            if self._on_status:
                self._on_status(f"⚠️ {nik}: Error process - {str(ex)}", False)
            return False, popup_action, status, saldo_jht, status_jmo
        finally:
            self._click_empty_area_to_close_popup()

    def _save_to_temp_list(self, row: int, kpj: str, nama: str, nik: str, status: str, popup_action: str, saldo_jht: str, status_jmo: str):
        """Simpan ke temporary list"""
        self.processed_rows.append({
            "RowNumber": row,
            "KPJ": kpj,
            "Nama": nama,
            "NIK": nik,
            "SaldoJHT": saldo_jht,
            "StatusJMO": status_jmo,
        })

    def _save_to_excel(self, output_path: str):
        """Simpan ke Excel (hanya yang punya saldo)"""
        try:
            # Baca file input
            input_wb = load_workbook(self.excel_path)
            input_sheet = input_wb.active
            
            total_cols = input_sheet.max_column
            
            # Buat file output
            output_wb = Workbook()
            output_sheet = output_wb.active
            output_sheet.title = "HASIL"
            
            # Copy header
            for col in range(1, total_cols + 1):
                output_sheet.cell(1, col, input_sheet.cell(1, col).value)
            
            # Tambah kolom Saldo JHT
            output_sheet.cell(1, total_cols + 1, "SALDO JHT")
            
            output_row = 2
            
            # Filter data dengan saldo > Rp 0
            data_dengan_saldo = [x for x in self.processed_rows if x["SaldoJHT"] != "Rp 0"]
            
            if self._on_status:
                self._on_status(f"📊 Menyimpan {len(data_dengan_saldo)} data dengan SALDO (dari total {len(self.processed_rows)})", True)
            
            for row_data in data_dengan_saldo:
                # Copy data asli
                for col in range(1, total_cols + 1):
                    output_sheet.cell(output_row, col, input_sheet.cell(row_data["RowNumber"], col).value)
                
                # Tambah Saldo JHT
                output_sheet.cell(output_row, total_cols + 1, row_data["SaldoJHT"])
                output_row += 1
            
            output_wb.save(output_path)
            
            # Tampilkan ringkasan
            print("\n" + "=" * 70)
            print("📊 RINGKASAN HASIL:")
            print("-" * 70)
            print(f"✅ Data dengan SALDO: {len(data_dengan_saldo)} data (DISIMPAN)")
            print(f"❌ Data tanpa saldo: {len(self.processed_rows) - len(data_dengan_saldo)} data (TIDAK DISIMPAN)")
            print(f"📊 TOTAL DIPROSES: {len(self.processed_rows)} data")
            print("-" * 70)
            
            if data_dengan_saldo:
                print("\n📋 DETAIL DATA YANG DISIMPAN:")
                for row in data_dengan_saldo[:10]:  # Tampilkan maksimal 10 data
                    print(f"✅ {row['NIK']} - {row['Nama']}: {row['SaldoJHT']}")
                if len(data_dengan_saldo) > 10:
                    print(f"... dan {len(data_dengan_saldo) - 10} data lainnya")
            
            print("=" * 70 + "\n")
            
            if self._on_status:
                self._on_status(f"💾 Excel berhasil disimpan: {output_path}", True)
            
        except Exception as ex:
            error_msg = f"Error save Excel: {str(ex)}"
            self._log(error_msg, "error")
            if self._on_status:
                self._on_status(f"❌ {error_msg}", False)

    def _show_console_confirmation(self, message: str) -> bool:
        """Tampilkan konfirmasi di console"""
        try:
            print("\n" + "█" * 80)
            print("██" + " " * 76 + "██")
            
            lines = message.split("\n")
            for line in lines:
                centered_line = line.ljust(76)[:76]
                print(f"██ {centered_line} ██")
            
            print("██" + " " * 76 + "██")
            print("██ " + "TEKAN ENTER UNTUK MELANJUTKAN...".ljust(76) + " ██")
            print("██" + " " * 76 + "██")
            print("█" * 80 + "\n")
            
            input()
            if self._on_status:
                self._on_status("▶ MELANJUTKAN proses...", True)
            return True
        except Exception as ex:
            if self._on_status:
                self._on_status(f"⚠️ Error menunggu konfirmasi: {str(ex)}", False)
            return False

    def stop(self):
        """Stop proses"""
        self.is_running = False
        if self._on_status:
            self._on_status("⏹ Menghentikan proses...", True)

    def start(self):
        """Start proses"""
        if not self.excel_path or not os.path.exists(self.excel_path):
            if self._on_status:
                self._on_status("❌ File Excel tidak ditemukan!", False)
            return
        
        self.is_running = True
        self.processed_rows.clear()
        self.sukses_count = 0
        self.gagal_count = 0
        
        self.current_output_path = os.path.join(
            os.path.dirname(self.excel_path),
            f"HASIL_JMO_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        )
        
        try:
            # Baca file Excel
            input_wb = load_workbook(self.excel_path)
            input_sheet = input_wb.active
            
            last_row = input_sheet.max_row
            
            if self._on_status:
                self._on_status(f"📊 Memproses {last_row - 1} data...", True)
            
            print("\n" + "=" * 70)
            print("🚀 BOT LASIK - DETEKSI STATUS JMO")
            print("✅ HANYA POPUP KONFIRMASI PERSETUJUAN YANG DIPROSES")
            print("❌ POPUP LAINNYA LANGSUNG DITUTUP (KLIK AREA KOSONG)")
            print("=" * 70 + "\n")
            
            for row in range(2, last_row + 1):
                if not self.is_running:
                    break
                
                kpj = str(input_sheet.cell(row, 1).value or "").strip()
                nama = str(input_sheet.cell(row, 2).value or "").strip()
                nik = str(input_sheet.cell(row, 3).value or "").strip()
                kelurahan = str(input_sheet.cell(row, 4).value or "").strip()
                
                if not nik or not kpj:
                    self.gagal_count += 1
                    continue
                
                if self._on_progress:
                    self._on_progress(row - 1, last_row - 1)
                
                self._log(f"Processing {row - 1}/{last_row - 1} - {nik}")
                
                try:
                    success, popup_action, status, saldo_jht, status_jmo = self._process_row(
                        row, kpj, nama, nik, kelurahan
                    )
                    
                    self._save_to_temp_list(row, kpj, nama, nik, status, popup_action, saldo_jht, status_jmo)
                    
                    if success:
                        self.sukses_count += 1
                        self._log(f"{nik} | {status_jmo} | {saldo_jht}", "success")
                    else:
                        self.gagal_count += 1
                        self._log(f"{nik} - {status_jmo}", "warning" if "SKIP" in status_jmo else "error")
                    
                    # Tampilkan konfirmasi di console
                    if row < last_row and not status_jmo.startswith("SKIP") and not status_jmo.startswith("GAGAL"):
                        message = f"📊 DATA KE-{row - 1} DARI {last_row - 1}\n" \
                                  f"━━━━━━━━━━━━━━━━━━━━\n" \
                                  f"👤 NIK       : {nik}\n" \
                                  f"👤 NAMA      : {nama}\n" \
                                  f"📱 KPJ       : {kpj}\n" \
                                  f"━━━━━━━━━━━━━━━━━━━━\n" \
                                  f"💰 SALDO JHT : {saldo_jht}\n" \
                                  f"━━━━━━━━━━━━━━━━━━━━"
                        
                        self._show_console_confirmation(message)
                    
                    self._clear_form_fields()
                    time.sleep(0.9)
                    
                except Exception as ex:
                    self.gagal_count += 1
                    if self._on_status:
                        self._on_status(f"⚠️ Error: {nik} - {str(ex)}", False)
                    self._save_to_temp_list(row, kpj, nama, nik, "ERROR", "EXCEPTION", "Rp 0", "ERROR SISTEM")
                    
                    try:
                        if self.driver:
                            self.driver.refresh()
                            time.sleep(0.8)
                    except:
                        pass
            
            if self.processed_rows:
                self._save_to_excel(self.current_output_path)
                print("\n" + "=" * 70)
                self._log(f"SELESAI | Sukses: {self.sukses_count} | Gagal: {self.gagal_count}", "success")
                print(f"💰 Total data dengan saldo: {len([x for x in self.processed_rows if x['SaldoJHT'] != 'Rp 0'])}")
                print("=" * 70 + "\n")
                if self._on_status:
                    self._on_status(f"🎯 SELESAI: {self.sukses_count} sukses, {self.gagal_count} gagal", True)
            else:
                if self._on_status:
                    self._on_status("❌ Tidak ada data yang berhasil diproses", False)
                
        except Exception as ex:
            error_msg = f"SYSTEM ERROR: {str(ex)}"
            self._log(error_msg, "error")
            if self._on_status:
                self._on_status(f"❌ {error_msg}", False)
        finally:
            try:
                if self.processed_rows:
                    self._save_to_excel(self.current_output_path)
                    if self._on_status:
                        self._on_status("💾 Data otomatis disimpan saat STOP / selesai", True)
            except Exception as ex:
                if self._on_status:
                    self._on_status(f"❌ Gagal save akhir: {str(ex)}", False)
            
            try:
                if self.driver:
                    self.driver.quit()
            except:
                pass
            
            self.is_running = False
